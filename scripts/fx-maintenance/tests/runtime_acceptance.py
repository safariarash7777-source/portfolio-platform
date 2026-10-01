"""Offline runtime acceptance. Synthetic fixtures stay outside Git."""
import atexit,base64,datetime as dt,hashlib,json,os,random,runpy,socket,subprocess,sys,time
from pathlib import Path
from unittest.mock import patch
import openpyxl,requests
from streamlit.testing.v1 import AppTest

ROOT=Path(__file__).resolve().parents[1]
D=ROOT/'work/liara-deploy-20260929/fx-dashboard'
FIX=Path(os.environ['FX_TEST_FIXTURES']).resolve();FIX.mkdir(parents=True,exist_ok=True)
sys.path.insert(0,str(D));sys.path.insert(0,str(ROOT/'work'))
os.environ.update(FX_DATA_DIR=str(FIX/'health'),FX_REQUIRE_AUTH='1')
import data_sources as ds
health=Path(os.environ['FX_DATA_DIR']);health.mkdir(exist_ok=True)

def digest(path):return hashlib.sha256(path.read_bytes()).hexdigest()
def write_status():
    now=dt.datetime.now(dt.timezone.utc)
    body={'run_id':'synthetic-runtime','status':'attention_required','execution_result':'completed','source_health':'degraded','data_completeness':'incomplete','updated_at':now.isoformat(),'finished_at':now.isoformat(),'cadence_days':5,'schedule':{'next_run_at':(now+dt.timedelta(days=5)).isoformat()}}
    raw=json.dumps(body).encode();(health/'agent-status.json').write_bytes(raw)
    (health/'agent-status-publication.json').write_text(json.dumps({'run_id':body['run_id'],'status':'verified','sha256':hashlib.sha256(raw).hexdigest()}),encoding='utf-8')
    return body,raw

def fixtures():
    rng=random.Random(732)
    book=openpyxl.Workbook();sheet=book.active;sheet.title='Data'
    for i,year in enumerate(range(1340,1405)):
        col=ds._FIRST_DATA_COL+1+i;sheet.cell(2,col,year+621);sheet.cell(3,col,year)
        for row,key in ds._ROW_MAP.items():
            value=(1.055**i)*(100+rng.random()*25) if not ('growth' in key or 'infl' in key) else 3+rng.random()*8
            sheet.cell(row+1,col,value)
    annual=FIX/'annual-synthetic.xlsx';book.save(annual)
    book=openpyxl.Workbook();sheet=book.active
    sheet.append(['دوره (سال/ماه)','تورم ماهانه (درصد)','تورم نقطه به نقطه (درصد)','تورم سالانه (درصد)'])
    for year in range(1399,1406):
        for month in ds.PERSIAN_MONTHS:sheet.append([f'{month} {year}',rng.uniform(0.2,2),rng.uniform(8,18),rng.uniform(8,18)])
    inflation=FIX/'inflation-synthetic.xlsx';book.save(inflation)
    book=openpyxl.Workbook();sheet=book.active
    for row in ['Synthetic benchmark','Not observed history','Fixture only']:sheet.append([row])
    sheet.append(['تاریخ میلادی','تاریخ شمسی','مقدار'])
    for year in range(1399,1406):
        for month in range(1,13):sheet.append([dt.datetime(*ds._jalali_to_gregorian(year,month,1)),f'{year}-{month:02d}-01',rng.uniform(.1,.2)])
    ytm=FIX/'ytm-synthetic.xlsx';book.save(ytm)
    # Only workbook column headers are read from the supplied schema source.
    schema=Path(os.environ['FX_TEST_BOURSE_SCHEMA'])
    original=openpyxl.load_workbook(schema,read_only=True,data_only=True)
    book=openpyxl.Workbook();book.remove(book.active)
    for name in ['صنایع','سهام','سکتورها']:
        headers=list(next(original[name].iter_rows(values_only=True)))
        sheet=book.create_sheet(name);sheet.append(headers)
        for i in range(8):
            sheet.append([f'آزمایشی {i+1}' if str(h) in ['صنعت','نماد','نام','سکتور','گروه','سناریو'] else rng.uniform(1,20) for h in headers])
    original.close();bourse=FIX/'bourse-synthetic.xlsx';book.save(bourse)
    os.environ.update(PPP_EXCEL=str(annual),INFLATION_XLSX=str(inflation),RISKFREE_YTM_XLSX=str(ytm),BOURSE_XLSX=str(bourse))
    # ds was imported to read schema; update paths before the AppTest import uses it.
    ds.EXCEL_PATH,ds.INFLATION_PATH,ds.YTM_PATH,ds.BOURSE_PATH=map(str,[annual,inflation,ytm,bourse])
    return [annual,inflation,ytm,bourse]

def run():
    files=fixtures();before={f.name:digest(f) for f in files};body,raw=write_status()
    from maintenance_status import view
    from auth_gate import verify_token,_sign
    import ifb_bonds,publish_agent_status
    result={};blocked=[]
    def deny(*a,**k):blocked.append('denied');raise requests.ConnectionError('Offline runtime acceptance')
    original_connect=socket.socket.connect
    original_connection=socket.create_connection
    def guard_connect(sock,address):
        if isinstance(address,tuple) and address[0] in ('127.0.0.1','::1','localhost'):
            return original_connect(sock,address)
        raise OSError('Upstream sockets disabled')
    def guard_connection(address,*args,**kwargs):
        if address[0] in ('127.0.0.1','::1','localhost'):return original_connection(address,*args,**kwargs)
        raise OSError('Upstream sockets disabled')
    salt=os.urandom(16);password=base64.urlsafe_b64encode(os.urandom(24)).decode()
    encoded='pbkdf2_sha256$300000$'+base64.urlsafe_b64encode(salt).decode()+'$'+base64.urlsafe_b64encode(hashlib.pbkdf2_hmac('sha256',password.encode(),salt,300000)).decode()
    os.environ['FX_STANDALONE_PASSWORD_HASH']=encoded;os.environ['FX_EMBED_SECRET']=''
    with patch.object(requests.Session,'request',side_effect=deny),patch('urllib.request.urlopen',side_effect=deny),patch.object(socket.socket,'connect',guard_connect),patch.object(socket,'create_connection',guard_connection):
        at=AppTest.from_file(str(D/'app.py'),default_timeout=240).run()
        assert not at.exception and len(at.tabs)==0 and len(at.text_input)==1
        result['anonymous_denied']=True
        at.text_input[0].set_value('invalid-fixture-password');at.button[0].click().run()
        assert len(at.tabs)==0 and not at.exception;result['wrong_password_denied']=True
        at.text_input[0].set_value(password);at.button[0].click().run()
        assert not at.exception, 'Authenticated render exception; details suppressed'
        assert len(at.tabs)==12
        result['authenticated_tabs']=[tab.label for tab in at.tabs]
        values=[str(getattr(w,'value','')) for w in at.markdown]
        for label in ['نتیجهٔ اجرا:','سلامت منابع:','کامل‌بودن داده‌ها:']:assert any(label in v for v in values)
        result['three_status_axes']=True
        button=next(b for b in at.button if b.key=='refresh_agent_status');button.click().run()
        assert not at.exception and len(at.tabs)==12;result['status_refresh_rerender']=True
        # Source outage must preserve the pointer bytes and original date.
        folder=health/'ifb-treasury';folder.mkdir(exist_ok=True);pointer=folder/'latest.json'
        pointer.write_text(json.dumps({'last_trade_date':'2026-09-30','snapshot_sha256':'synthetic-marker'}),encoding='utf-8');previous=pointer.read_bytes()
        with patch.object(ifb_bonds,'DATA_DIR',health),patch.object(ifb_bonds.time,'sleep'):
            state=ifb_bonds.collect()
        assert pointer.read_bytes()==previous and state['previous_snapshot_retained'] and len(state['attempts'])==2
        result['outage_preserves_last_valid_pointer']=True
    receipt=json.loads((health/'agent-status-publication.json').read_text(encoding='utf-8'))
    assert view(body,raw,receipt)['publication_verified']
    assert not view(body,raw+b' ',receipt)['publication_verified']
    assert view(body,raw,receipt,now=dt.datetime.now(dt.timezone.utc)+dt.timedelta(days=6))['overdue']
    result['checksum_and_overdue_rules']=True
    # Verify the runbook junction and archived source definitions without fetching.
    assert (D/'.data_health').resolve()==health.resolve()
    from validate_cbi_mapping import read
    entries=json.loads((health/'cbi-tsd-health.json').read_text(encoding='utf-8'))['series']
    for entry in entries:read(entry)
    result['archived_CBI_workbooks_validated']=len(entries)
    result['data_health_junction_verified']=True
    import cbi_tsd_catalog
    assert cbi_tsd_catalog.ROOT==health.resolve()
    marker=b'Offline portability fixture; no financial observations'
    archived=cbi_tsd_catalog.archive(marker,entries[0]['id'],'.fixture')
    assert (health/'cbi-tsd'/entries[0]['id']/(archived+'.fixture')).read_bytes()==marker
    result['CBI_archiver_long_path_write']=True
    # Exercise the actual publisher protocol with local transport/readback, never SSH.
    remote=FIX/'simulated-remote';remote.mkdir(exist_ok=True)
    def local_transport(args,**kwargs):
        assert args[2]=='--upload-verified'
        source=Path(args[3]);target=remote/Path(args[4]).name
        target.write_bytes(source.read_bytes());sha=digest(target);assert sha==digest(source)
        import subprocess
        return subprocess.CompletedProcess(args,0,json.dumps({'status':'verified','sha256':sha}),'')
    (ROOT/'outputs').mkdir(exist_ok=True)
    with patch.object(publish_agent_status.subprocess,'run',side_effect=local_transport):
        pub=publish_agent_status.publish_status(health/'agent-status.json','/isolated-test')
    assert pub['status']=='verified'
    assert digest(remote/'agent-status.json')==digest(health/'agent-status.json')
    received=json.loads((remote/'agent-status-publication.json').read_text(encoding='utf-8'))
    assert received['sha256']==digest(remote/'agent-status.json');result['publisher_simulated_transport_readback']=True
    # Run only template orchestration. Child fetch/publication scripts are stubbed,
    # and schedule comes from a fixture; no installed hook or real cron is touched.
    os.environ.update(FX_MAINTENANCE_ROOT=str(ROOT),FX_LOCAL_PYTHON=sys.executable)
    child_calls=[]
    def isolated_child(args,**kwargs):
        if len(args)>2 and args[2]=='--upload-verified':return local_transport(args,**kwargs)
        script=Path(args[1]);assert script.is_relative_to(ROOT) and script.exists()
        child_calls.append(script.name)
        return subprocess.CompletedProcess(args,0,json.dumps({'status':'failed','series':[],'error_type':'OfflineFixture'}),'')
    with patch.object(publish_agent_status,'schedule_metadata',return_value={'id':'fixture-only','next_run_at':body['schedule']['next_run_at']}),patch.object(subprocess,'run',side_effect=isolated_child):
        namespace=runpy.run_path(str(ROOT/'hermes/fx_economic_preflight.py'),run_name='fx_isolated_template')
    namespace['lock_file'].close();atexit.unregister(namespace['finalize_unexpected_exit'])
    final=json.loads((health/'agent-status.json').read_text(encoding='utf-8'))
    assert final['execution_result']=='completed_with_errors' and final['data_completeness']=='incomplete'
    assert final['schedule']['id']=='fixture-only' and len(child_calls)>=8
    assert digest(remote/'agent-status.json')==digest(health/'agent-status.json')
    result['template_orchestration_with_stubbed_children']=True
    result['template_child_scripts_executed_live']=False
    secret=base64.urlsafe_b64encode(os.urandom(24)).decode();exp=str(int(time.time())+60)
    assert verify_token(exp+'.fixture-admin.'+_sign(secret,exp+'.fixture-admin'),secret)[0]
    assert not verify_token('invalid',secret)[0];result['signed_token_validation']=True
    assert before=={f.name:digest(f) for f in files};result['fixture_inputs_unchanged']=True
    # Real Streamlit process with test-only socket and DNS restrictions.
    with socket.socket() as sock:
        sock.bind(('127.0.0.1',0));port=sock.getsockname()[1]
    env=dict(os.environ,PYTHONPATH=str(ROOT/'tests/network_guard'),STREAMLIT_BROWSER_GATHER_USAGE_STATS='false')
    with (FIX/'startup-private.log').open('w',encoding='utf-8') as log:
        process=subprocess.Popen([sys.executable,'-m','streamlit','run',str(D/'app.py'),'--server.address=127.0.0.1',f'--server.port={port}','--server.headless=true','--browser.gatherUsageStats=false'],env=env,cwd=FIX,stdout=log,stderr=log)
        try:
            deadline=time.monotonic()+35
            while time.monotonic()<deadline:
                assert process.poll() is None,'Startup process exited'
                try:
                    client=requests.Session();client.trust_env=False
                    health_reply=client.get(f'http://127.0.0.1:{port}/_stcore/health',timeout=2)
                    if health_reply.status_code==200:break
                except requests.RequestException:time.sleep(.25)
            else:raise AssertionError('Startup health timeout')
            assert client.get(f'http://127.0.0.1:{port}/',timeout=2).status_code==200
            result['real_streamlit_startup_loopback']=True
        finally:
            process.terminate();process.wait(timeout=15)
    result.update(status='PASS',upstream_connections=0,blocked_source_attempts=len(blocked),synthetic_data=True,financial_methodology_validated=False,live_hook_or_production_changed=False)
    (FIX/'runtime-result.json').write_text(json.dumps(result,ensure_ascii=False,indent=2),encoding='utf-8')
    print(json.dumps({k:v for k,v in result.items() if k!='authenticated_tabs'}))

if __name__=='__main__':run()
