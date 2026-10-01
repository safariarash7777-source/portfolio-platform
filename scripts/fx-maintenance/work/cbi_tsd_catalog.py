"""Official CBI TSD catalogue and Excel receipts; no observations enter model output."""
import datetime as dt, hashlib,json,sys
from pathlib import Path
import requests
from concurrent.futures import ThreadPoolExecutor
import openpyxl,io,math,re,shutil,warnings
warnings.filterwarnings("ignore",message="Workbook contains no default style")
PROJECT=Path(__file__).resolve().parent.parent
ROOT=(PROJECT/'work/liara-deploy-20260929/fx-dashboard/.data_health').resolve()
BASE='https://tsdview.cis.cbi.ir/backend/api/'
IDS=['3505fc25-6651-8f09-e063-4ceaa8c0b336','3505fc25-5fec-8f09-e063-4ceaa8c0b336','3505fc25-6648-8f09-e063-4ceaa8c0b336','3505fc25-665f-8f09-e063-4ceaa8c0b336','3505fc25-5fb9-8f09-e063-4ceaa8c0b336']
def session():
 s=requests.Session();s.trust_env=False;s.headers.update({'Accept-Language':'fa','Accept':'application/json'});return s
def archive(body,folder,suffix):
 digest=hashlib.sha256(body).hexdigest();p=ROOT/'cbi-tsd'/folder/(digest+suffix);p.parent.mkdir(parents=True,exist_ok=True)
 if not p.exists():p.write_bytes(body)
 return digest

def collect():
 report={'source':'https://tsdview.cis.cbi.ir/single-data','checked_at':dt.datetime.now(dt.timezone.utc).isoformat(),'status':'checking','series':[]}
 try:
  r=session().post(BASE+'category/',json={},timeout=(12,110));r.raise_for_status();rows=r.json()
  if not isinstance(rows,list) or not all(isinstance(x,dict) and 'id' in x and 'title' in x for x in rows):raise ValueError('catalog_schema')
  report['catalog_sha256']=archive(r.content,'catalog','.json');report['catalog_count']=len(rows)
 except (requests.RequestException,ValueError) as e:
  report.update(status='source_unavailable',error_type=type(e).__name__);return report
 catalog={x['id']:x for x in rows}
 def export(sid, monthly=False):
  x=catalog[sid];meta={'id':sid,'title':x['title'],'definition_path':x['path'],'unit':x['unit']['title'],'frequency':x['subFrequency']['title'],'catalog_first_year':x['startYear'],'catalog_last_year':x['endYear'],'calendar':x['calenderType'],'preliminary':x['preliminary'],'publication_date':None,'excel_binding':'pending_definition_and_period_mapping'}
  if monthly:
   meta['frequency']='ماهانه'
   meta['series_variant']='CPI_monthly'
  try:
   payload={'fromDate':str(x['startYear'])+'0101','toDateDate':str(x['endYear'])+'1230','categoryIds':[sid],'calender':x['calenderType'].upper(),'subFrequencyId':x['subFrequencyId']}
   if monthly:payload['subFrequencyId']='3B7F33D82FE0D7C5FC1CDA862B7F433514B738A86C5E246F9E0A0B1FADD68FBB'
   resp=session().post(BASE+'value/export-excel',json=payload,timeout=(12,110));resp.raise_for_status()
   if not resp.content.startswith(b'PK'):raise ValueError('not_xlsx')
   wb=openpyxl.load_workbook(io.BytesIO(resp.content),read_only=True,data_only=True)
   observations=[]
   for row in wb.active.iter_rows(min_row=9,values_only=True):
    if row[0] is not None and row[1] is not None:
     if not isinstance(row[1],(int,float)) or not math.isfinite(row[1]):raise ValueError('invalid_observation')
     observations.append(str(row[0]))
   if not observations:raise ValueError('empty_observations')
   meta.update(observation_count=len(observations),first_observed_period=observations[0],last_observed_period=observations[-1])
   sheets=[{'name':sh.title,'rows':sh.max_row,'columns':sh.max_column} for sh in wb.worksheets]
   meta.update(status='official_excel_received',sha256=archive(resp.content,sid,'.xlsx'),sheets=sheets)
   wb.close()
   output=PROJECT/'outputs/cbi-official-excel';output.mkdir(parents=True,exist_ok=True)
   original=ROOT/'cbi-tsd'/sid/(meta['sha256']+'.xlsx')
   shutil.copy2(original,output/(sid+('-monthly' if monthly else '')+'.xlsx'))
  except (requests.RequestException,ValueError,OSError) as e:meta.update(status='export_failed',error_type=type(e).__name__)
  return meta
 with ThreadPoolExecutor(max_workers=2) as pool:report['series']=list(pool.map(export,[i for i in IDS if i in catalog]))
 report['series'].append(export(IDS[0],monthly=True))
 report['status']='exports_received_mapping_pending' if any(x['status']=='official_excel_received' for x in report['series']) else 'catalog_received_exports_failed'
 return report
if __name__=='__main__':
 sys.stdout.reconfigure(encoding='utf-8');r=collect();ROOT.mkdir(parents=True,exist_ok=True)
 sys.path.insert(0,str(PROJECT/'work/liara-deploy-20260929/fx-dashboard'))
 from source_health import atomic_write
 previous=ROOT/'cbi-tsd-last-valid.json'
 if len(r.get('series',[]))==6 and all(s.get('status')=='official_excel_received' for s in r['series']):atomic_write(previous,r)
 else:r['previous_snapshot_retained']=previous.exists()
 atomic_write(ROOT/'cbi-tsd-health.json',r);print(json.dumps(r,ensure_ascii=False))
