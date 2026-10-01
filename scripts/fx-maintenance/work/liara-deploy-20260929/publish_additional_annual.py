"""Publish five definition-matched annual series with rollback and model gates."""
import hashlib
import json
import os
from pathlib import Path
import shutil
import subprocess
import sys
import tarfile
import datetime as dt

def digest(path):return hashlib.sha256(path.read_bytes()).hexdigest()

def remote(stage,rollback=False):
    import data_sources as ds
    from source_health import DATA_DIR
    target=Path(ds.EXCEL_PATH)
    if rollback:
        shutil.copyfile(stage/'previous.xlsx',target)
        if (stage/'previous-metadata.json').exists():
            shutil.copyfile(stage/'previous-metadata.json',DATA_DIR/'annual-economic-update.json')
        if (stage/'previous-gdp-metadata.json').exists():shutil.copyfile(stage/'previous-gdp-metadata.json',DATA_DIR/'annual-update-latest.json')
        print(json.dumps({'status':'rolled_back'}));return
    metadata=json.loads((stage/'annual-economic-update.json').read_text(encoding='utf-8'))
    candidate=stage/'candidate.xlsx'
    assert digest(candidate)==metadata['after_sha256']
    if digest(target)==metadata['after_sha256']:
        print(json.dumps({'status':'already_published'}));return
    assert digest(target)==metadata['before_sha256'],'Concurrent annual changes; refusing overwrite'
    import openpyxl
    current=openpyxl.load_workbook(target,data_only=True)['Data']
    proposed=openpyxl.load_workbook(candidate,data_only=True)['Data']
    assert current.max_row==proposed.max_row and current.max_column==proposed.max_column
    changes={(r,c) for r in range(1,current.max_row+1) for c in range(1,current.max_column+1)
             if current.cell(r,c).value!=proposed.cell(r,c).value}
    allowed={(r,c) for r in metadata['allowed_rows'] for c in range(4,proposed.max_column+1)
             if isinstance(proposed.cell(2,c).value,(int,float)) and proposed.cell(2,c).value<=metadata['allowed_calendar_max']}
    assert changes and changes.issubset(allowed)
    backup=DATA_DIR/'backups'/('annual-'+stage.name);backup.mkdir(parents=True,exist_ok=False)
    shutil.copyfile(target,backup/'previous.xlsx')
    (backup/'checksums.json').write_text(json.dumps({'previous.xlsx':digest(target),'candidate.xlsx':digest(candidate)}))
    previous=DATA_DIR/'annual-economic-update.json'
    shutil.copyfile(DATA_DIR/'annual-update-latest.json',backup/'previous-gdp-metadata.json')
    if previous.exists():shutil.copyfile(previous,backup/'previous-metadata.json')
    # Keep source vintages and metadata before switching the workbook atomically.
    for item in (stage/'official-documents').rglob('*'):
        if item.is_file():
            destination=DATA_DIR/'official-documents'/item.relative_to(stage/'official-documents')
            destination.parent.mkdir(parents=True,exist_ok=True);shutil.copyfile(item,destination)
    for name in ['annual-economic-update.json','worldbank-additional-check.json','annual-update-latest.json']:
        temporary=DATA_DIR/(name+'.pending');shutil.copyfile(stage/name,temporary);os.replace(temporary,DATA_DIR/name)
    temporary=target.with_name('annual.pending.xlsx');shutil.copyfile(candidate,temporary);os.replace(temporary,target)
    print(json.dumps({'status':'annual_published','after_sha256':digest(target),'backup':str(backup)}))

def local():
    root=Path(__file__).parent;sys.path.insert(0,str(root/'fx-dashboard'))
    import data_sources as ds
    from source_health import DATA_DIR
    metadata=json.loads((DATA_DIR/'annual-economic-update.json').read_text(encoding='utf-8'))
    assert digest(Path(ds.EXCEL_PATH))==metadata['after_sha256']
    runtime=os.environ['FX_BUNDLED_PYTHON']
    def ssh(*args):
        result=subprocess.run([runtime,str(root/'vm_command.py'),*args],capture_output=True,text=True,encoding='utf-8',timeout=300)
        if result.returncode:raise RuntimeError('Annual publication operation failed')
        return result.stdout.strip()
    ssh('--upload',str(root/'verify_us_version.py'),'/tmp/verify_us_version.py')
    current=json.loads(ssh('docker cp /tmp/verify_us_version.py fx-dashboard:/app/verify_us_version.py && docker exec -w /app fx-dashboard python verify_us_version.py'))
    if current['data'][Path(ds.EXCEL_PATH).name]==metadata['after_sha256']:
        print(json.dumps({'status':'already_published','official_period':metadata['allowed_calendar_max']}));return
    stamp=dt.datetime.now(dt.timezone.utc).strftime('%Y%m%dT%H%M%S%fZ')
    package=root/('annual-'+stamp+'.tar.gz')
    with tarfile.open(package,'w:gz') as archive:
        archive.add(ds.EXCEL_PATH,arcname='candidate.xlsx')
        for name in ['annual-economic-update.json','worldbank-additional-check.json','annual-update-latest.json']:archive.add(DATA_DIR/name,arcname=name)
        for folder in (DATA_DIR/'official-documents').glob('WB_*'):archive.add(folder,arcname='official-documents/'+folder.name)
    archive_path='/tmp/annual-'+stamp+'.tar.gz'
    ssh('--upload',str(package),archive_path)
    ssh('--upload',str(Path(__file__)),'/tmp/publish_additional_annual.py')
    stage='/opt/fx-dashboard/data-health/staging/'+stamp
    ssh(f'mkdir -p {stage} && tar -xzf {archive_path} -C {stage} && docker cp /tmp/publish_additional_annual.py fx-dashboard:/app/publish_additional_annual.py')
    gate=ssh(f'docker exec -e FX_REQUIRE_AUTH=0 -e PPP_EXCEL=/data-health/staging/{stamp}/candidate.xlsx -w /app fx-dashboard python precise_model_gate.py')
    assert '"exceptions": []' in gate and '"tabs": 12' in gate
    result=json.loads(ssh(f'docker exec -w /app fx-dashboard python publish_additional_annual.py --remote /data-health/staging/{stamp}'))
    try:
        # Password is handled by the SSH helper, never by the maintenance model.
        smoke=ssh('--smoke-auth')
        assert 'DASHBOARD_TAB_COUNT 12' in smoke and 'AFTER_LOGIN_EXCEPTIONS []' in smoke
        final=json.loads(ssh('docker exec -w /app fx-dashboard python verify_us_version.py'))
        assert final['data'][Path(ds.EXCEL_PATH).name]==metadata['after_sha256']
        ssh('docker exec -w /app fx-dashboard python macro_health.py > /opt/fx-dashboard/data-health/macro-health-last-run.log && curl -fsS https://62.60.191.24/_stcore/health')
    except Exception:
        if result.get('backup'):
            ssh('docker exec -w /app fx-dashboard python publish_additional_annual.py --rollback '+result['backup'])
        raise
    print(json.dumps({'status':'annual_published_and_models_verified','official_period':metadata['allowed_calendar_max'],'after_sha256':metadata['after_sha256']}))

if __name__=='__main__':
    if '--remote' in sys.argv:remote(Path(sys.argv[-1]))
    elif '--rollback' in sys.argv:remote(Path(sys.argv[-1]),True)
    else:local()
