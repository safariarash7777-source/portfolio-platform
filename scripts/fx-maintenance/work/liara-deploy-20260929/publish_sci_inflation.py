"""Publish official SCI revisions with checksum conflict checks and a model gate."""
import datetime as dt
import hashlib,json,os,shutil,subprocess,sys,tarfile
from pathlib import Path

def sha(p):return hashlib.sha256(p.read_bytes()).hexdigest()

def remote(stage,rollback=False):
    import data_sources as ds
    from source_health import DATA_DIR
    target=Path(ds.INFLATION_PATH);receipt=DATA_DIR/'sci-inflation-update.json'
    if rollback:
        shutil.copyfile(stage/'previous.xlsx',target)
        if (stage/'previous.json').exists():shutil.copyfile(stage/'previous.json',receipt)
        print(json.dumps({'status':'rolled_back'}));return
    meta=json.loads((stage/'receipt.json').read_text(encoding='utf-8'))
    assert sha(stage/'candidate.xlsx')==meta['after_sha256']
    if sha(target)==meta['after_sha256']:
        print(json.dumps({'status':'already_published'}));return
    assert sha(target)==meta['before_sha256'],'Concurrent inflation revision; refusing overwrite'
    import openpyxl
    a=openpyxl.load_workbook(target,data_only=True).active;b=openpyxl.load_workbook(stage/'candidate.xlsx',data_only=True).active
    assert a.max_column==b.max_column and b.max_row>=a.max_row
    allowed=set(meta['updated_periods'])
    for r in range(1,b.max_row+1):
        period=b.cell(r,1).value
        if r<=a.max_row:assert a.cell(r,1).value==period
        for c in range(1,b.max_column+1):
            old=a.cell(r,c).value if r<=a.max_row else None
            if old!=b.cell(r,c).value:assert period in allowed and (c in (2,3,4,5) or (r>a.max_row and c==1))
    backup=DATA_DIR/'backups'/('sci-'+stage.name);backup.mkdir(parents=True,exist_ok=False)
    shutil.copyfile(target,backup/'previous.xlsx')
    if receipt.exists():shutil.copyfile(receipt,backup/'previous.json')
    (backup/'checksums.json').write_text(json.dumps({'before_sha256':sha(target),'after_sha256':sha(stage/'candidate.xlsx')}))
    for source in (stage/'documents').glob('*.pdf'):
        out=DATA_DIR/'sci-validated'/source.name;out.parent.mkdir(parents=True,exist_ok=True);shutil.copyfile(source,out)
    temporary=target.with_name('inflation.pending.xlsx');shutil.copyfile(stage/'candidate.xlsx',temporary);os.replace(temporary,target)
    temporary=receipt.with_suffix('.pending');shutil.copyfile(stage/'receipt.json',temporary);os.replace(temporary,receipt)
    print(json.dumps({'status':'published','backup':str(backup),'after_sha256':sha(target)}))

def local():
    root=Path(__file__).parent;sys.path.insert(0,str(root/'fx-dashboard'))
    import data_sources as ds
    from source_health import DATA_DIR
    meta=json.loads((DATA_DIR/'sci-inflation-update.json').read_text(encoding='utf-8'));assert sha(Path(ds.INFLATION_PATH))==meta['after_sha256']
    runtime=os.environ['FX_BUNDLED_PYTHON']
    def ssh(*args):
        result=subprocess.run([runtime,str(root/'vm_command.py'),*args],capture_output=True,text=True,encoding='utf-8',timeout=360)
        if result.returncode:raise RuntimeError('SCI publication operation failed')
        return result.stdout.strip()
    current=json.loads(ssh("docker exec fx-dashboard python -c \"import os,hashlib,json;from pathlib import Path;p=Path(os.environ['INFLATION_XLSX']);print(json.dumps({'sha256':hashlib.sha256(p.read_bytes()).hexdigest()}))\""))
    if current['sha256']==meta['after_sha256']:
        print(json.dumps({'status':'already_published','last_period':meta['last_period'],'after_sha256':current['sha256']}));return
    assert current['sha256']==meta['before_sha256'],'Concurrent revision; refusing overwrite'
    stamp=dt.datetime.now(dt.timezone.utc).strftime('%Y%m%dT%H%M%S%fZ');package=root/('sci-'+stamp+'.tar.gz')
    with tarfile.open(package,'w:gz') as t:
        t.add(ds.INFLATION_PATH,arcname='candidate.xlsx');t.add(DATA_DIR/'sci-inflation-update.json',arcname='receipt.json')
        for pdf in (DATA_DIR/'sci-validated').glob('*.pdf'):t.add(pdf,arcname='documents/'+pdf.name)
    stage='/opt/fx-dashboard/data-health/staging/sci-'+stamp;inside='/data-health/staging/sci-'+stamp
    ssh('--upload',str(package),'/tmp/sci-'+stamp+'.tar.gz');ssh('--upload',str(Path(__file__)),'/tmp/publish_sci_inflation.py')
    ssh(f'mkdir -p {stage} && tar -xzf /tmp/sci-{stamp}.tar.gz -C {stage} && docker cp /tmp/publish_sci_inflation.py fx-dashboard:/app/publish_sci_inflation.py')
    # Test a candidate in a separate process; the live process retains auth settings.
    gate=ssh(f'docker exec -e FX_REQUIRE_AUTH=0 -e INFLATION_XLSX={inside}/candidate.xlsx -w /app fx-dashboard python precise_model_gate.py')
    assert '"exceptions": []' in gate and '"tabs": 12' in gate
    result=json.loads(ssh(f'docker exec -w /app fx-dashboard python publish_sci_inflation.py --remote {inside}'))
    try:
        smoke=ssh('--smoke-auth');assert 'DASHBOARD_TAB_COUNT 12' in smoke and 'AFTER_LOGIN_EXCEPTIONS []' in smoke
        ssh('curl -fsS https://62.60.191.24/_stcore/health')
    except Exception:
        if result.get('backup'):ssh('docker exec -w /app fx-dashboard python publish_sci_inflation.py --rollback '+result['backup'])
        raise
    print(json.dumps({'status':'sci_published_and_models_verified','last_period':meta['last_period'],'after_sha256':meta['after_sha256']}))

if __name__=='__main__':
    if '--remote' in sys.argv:remote(Path(sys.argv[-1]))
    elif '--rollback' in sys.argv:remote(Path(sys.argv[-1]),True)
    else:local()
