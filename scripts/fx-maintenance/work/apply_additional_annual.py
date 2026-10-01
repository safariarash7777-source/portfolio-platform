from pathlib import Path
import sys,json,math,hashlib,shutil,datetime as dt,zipfile
from lxml import etree as ET
import openpyxl
sys.path.insert(0,'work/liara-deploy-20260929/fx-dashboard');import data_sources as ds
from source_health import DATA_DIR
report=json.loads((DATA_DIR/'worldbank-additional-check.json').read_text(encoding='utf-8'));entries={e['series']:e for e in report['series']};assert all(e['status']=='definition_and_snapshot_validated' for e in entries.values())
rows={'gdp_usd':11,'usa_cpi':13,'usa_infl':14,'usa_m2':15,'usa_m2_growth':16};divisors={'gdp_usd':1e9,'usa_m2':1e9,'usa_cpi':1,'usa_infl':1,'usa_m2_growth':1};observations={}
for name,e in entries.items():
 p=Path(e['document']);assert hashlib.sha256(p.read_bytes()).hexdigest()==e['sha256'];observations[name]={int(t['date']):t['value'] for t in json.loads(p.read_text(encoding='utf-8'))[1]}
for y,v in observations['usa_m2'].items():
 prev=observations['usa_m2'].get(y-1);gr=observations['usa_m2_growth'].get(y)
 if v is not None and prev is not None and gr is not None:assert abs((v/prev-1)*100-gr)<0.05
assert len(entries['gdp_usd']['legacy_matched_periods'])>=10, 'Historical billion-dollar scale must match'
f=Path(ds.EXCEL_PATH);before=hashlib.sha256(f.read_bytes()).hexdigest();w=openpyxl.load_workbook(f,data_only=True);s=w['Data'];changes=[]
for name,row in rows.items():
 for c in range(6,s.max_column+1):
  year=s.cell(2,c).value
  if not isinstance(year,(int,float)) or year>dt.datetime.now(dt.timezone.utc).year-1:continue
  year=int(year);v=observations[name].get(year);v=None if v is None else v/divisors[name];old=s.cell(row,c).value
  equal=old==v or (isinstance(old,(int,float)) and isinstance(v,(int,float)) and math.isclose(old,v,rel_tol=0,abs_tol=1e-10))
  if not equal:changes.append({'series':name,'row':row,'column':c,'period':year,'old_value':old,'new_value':v,'source_sha256':entries[name]['sha256']})
if not changes:
 print(json.dumps({'status':'already_current','after_sha256':before}));sys.exit(0)
stamp=dt.datetime.now(dt.timezone.utc).strftime('%Y%m%dT%H%M%S%fZ');folder=DATA_DIR/'economic-revisions'/stamp;folder.mkdir(parents=True);shutil.copy2(f,folder/f.name)
ns={'x':'http://schemas.openxmlformats.org/spreadsheetml/2006/main'}
with zipfile.ZipFile(f) as z:
 wb=ET.fromstring(z.read('xl/workbook.xml'));node=next(t for t in wb.findall('x:sheets/x:sheet',ns) if t.get('name')=='Data');rid=node.get('{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id');rels=ET.fromstring(z.read('xl/_rels/workbook.xml.rels'));target=next(t.get('Target') for t in rels if t.get('Id')==rid);xmlpath=target.lstrip('/') if target.startswith('/') else 'xl/'+target;tree=ET.fromstring(z.read(xmlpath))
 for ch in changes:
  addr=openpyxl.utils.get_column_letter(ch['column'])+str(ch['row']);cell=next(t for t in tree.findall('.//x:c',ns) if t.get('r')==addr)
  for child in list(cell):
   if child.tag.endswith(('}v','}f','}is')):cell.remove(child)
  cell.set('t','n')
  if ch['new_value'] is not None:ET.SubElement(cell,'{'+ns['x']+'}v').text=repr(ch['new_value'])
 temp=f.with_suffix('.pending.xlsx')
 with zipfile.ZipFile(temp,'w',compression=zipfile.ZIP_DEFLATED) as out:
  for member in z.infolist():out.writestr(member,ET.tostring(tree,encoding='utf-8',xml_declaration=True) if member.filename==xmlpath else z.read(member.filename))
test=openpyxl.load_workbook(temp,data_only=True)['Data'];actual={(r,c) for r in range(1,s.max_row+1) for c in range(1,s.max_column+1) if s.cell(r,c).value!=test.cell(r,c).value};assert actual=={(ch['row'],ch['column']) for ch in changes};temp.replace(f)
(folder/'annual-revisions-raw.json').write_text(json.dumps(changes),encoding='utf-8');meta={'status':'matching_annual_series_updated','changed_at':dt.datetime.now(dt.timezone.utc).isoformat(),'file':f.name,'before_sha256':before,'after_sha256':hashlib.sha256(f.read_bytes()).hexdigest(),'backup':str(folder/f.name),'changed_cell_count':len(changes),'allowed_rows':list(rows.values()),'allowed_calendar_max':dt.datetime.now(dt.timezone.utc).year-1,'series':[{'series':name,'last_period':e['last_period'],'definition':e['definition'],'calendar':'Gregorian','unit_divisor':divisors[name],'source':e['source'],'source_sha256':e['sha256'],'source_last_updated':e['source_last_updated'],'publication_date':None} for name,e in entries.items()],'checks':['Historical unit matching','USA money level/growth agreement','future columns preserved','all other cells preserved'],'missing_official_observations_kept_null':True};(DATA_DIR/'annual-economic-update.json').write_text(json.dumps(meta,ensure_ascii=False),encoding='utf-8');print(json.dumps(meta,ensure_ascii=True))

gdp_path=DATA_DIR/'annual-update-latest.json'
if gdp_path.exists():
 gdp=json.loads(gdp_path.read_text(encoding='utf-8'));gdp['after_sha256']=meta['after_sha256'];gdp['other_series_revision']='annual-economic-update.json';gdp_path.write_text(json.dumps(gdp,ensure_ascii=False),encoding='utf-8')
